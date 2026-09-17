#!/usr/bin/env python3
"""
ONE Slide Builder — ONEブランドデザインに準拠したpptx生成ヘルパー

使用方法:
    python3 one_slide_builder.py --template "Theme _ ONE for Business 2025.pptx" --output output.pptx

このスクリプトはClaude Desktopスキルから呼び出され、
ONEブランドデザインシステムに完全準拠したスライドを生成する。
"""

import argparse
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData


# ============================================================
# ONE Design Tokens
# ============================================================
ONE_GREEN = RGBColor(0x3E, 0xE3, 0xA6)
ONE_DARK_GREEN = RGBColor(0x00, 0xD1, 0x84)
ONE_LIGHT_GREEN = RGBColor(0xD7, 0xFF, 0xF0)
BLACK = RGBColor(0x00, 0x00, 0x00)
ONE_LIGHT_GRAY = RGBColor(0xF0, 0xF3, 0xF6)
ONE_DARK_GRAY = RGBColor(0x8C, 0x90, 0x9B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

CHART_COLORS = [ONE_GREEN, ONE_DARK_GREEN, ONE_DARK_GRAY, ONE_LIGHT_GRAY, BLACK]

# スライドサイズ
SLIDE_WIDTH = Emu(9144000)
SLIDE_HEIGHT = Emu(5143500)

# マージン
MARGIN_LEFT = Emu(508000)
MARGIN_TOP = Emu(411480)
CONTENT_TOP = Emu(771525)
CONTENT_WIDTH = Emu(8128000)

# フォント
FONT_TITLE = "Noto Sans JP"
FONT_BODY = "Noto Sans JP"
FONT_KPI = "Arial Black"


# ============================================================
# ベースヘルパー
# ============================================================

def add_blank_slide(prs, bg_color=None):
    """ブランクスライドを追加"""
    layout = prs.slide_layouts[4]  # Blank
    slide = prs.slides.add_slide(layout)
    if bg_color:
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = bg_color
    return slide


def add_title(slide, text, font_size=Pt(28), bold=True, color=BLACK):
    """スライドタイトルを追加"""
    txBox = slide.shapes.add_textbox(
        MARGIN_LEFT, MARGIN_TOP, CONTENT_WIDTH, Emu(400000)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_TITLE
    p.font.size = font_size
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.LEFT
    return txBox


def add_green_accent_line(slide):
    """タイトル下のグリーンアクセントライン"""
    line = slide.shapes.add_connector(
        1, MARGIN_LEFT, Emu(650000),
        MARGIN_LEFT + Emu(1200000), Emu(650000)
    )
    line.line.color.rgb = ONE_DARK_GREEN
    line.line.width = Pt(3)
    return line


def add_footnote(slide, text):
    """ソース・脚注（14pt, #8C909B）"""
    txBox = slide.shapes.add_textbox(
        MARGIN_LEFT, Emu(4800000), CONTENT_WIDTH, Emu(250000)
    )
    p = txBox.text_frame.paragraphs[0]
    p.text = text
    p.font.name = FONT_BODY
    p.font.size = Pt(14)
    p.font.color.rgb = ONE_DARK_GRAY
    p.alignment = PP_ALIGN.RIGHT
    return txBox


# ============================================================
# パターン: ステートメント
# ============================================================

def build_statement_slide(prs, statement, sub_text=""):
    """ステートメントスライド（#F0F3F6 BG）"""
    slide = add_blank_slide(prs, bg_color=ONE_LIGHT_GRAY)
    txBox = slide.shapes.add_textbox(
        MARGIN_LEFT, Emu(1500000), CONTENT_WIDTH, Emu(800000)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = statement
    p.font.name = FONT_TITLE
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.alignment = PP_ALIGN.LEFT

    if sub_text:
        stx = slide.shapes.add_textbox(
            MARGIN_LEFT, Emu(2500000), CONTENT_WIDTH, Emu(600000)
        )
        sp = stx.text_frame.paragraphs[0]
        sp.text = sub_text
        sp.font.name = FONT_BODY
        sp.font.size = Pt(20)
        sp.font.color.rgb = ONE_DARK_GRAY
    return slide


# ============================================================
# パターン: KPI-Row
# ============================================================

def _add_kpi_card(slide, left, top, width, height,
                  number, label, unit="", delta="", card_bg=ONE_LIGHT_GRAY):
    """単一KPIカード"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = card_bg
    shape.line.fill.background()
    shape.adjustments[0] = 0.05

    # 数値
    txBox = slide.shapes.add_textbox(
        left + Emu(80000), top + Emu(200000),
        width - Emu(160000), Emu(600000)
    )
    tf = txBox.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = str(number)
    run.font.name = FONT_KPI
    run.font.size = Pt(48)
    run.font.bold = True
    run.font.color.rgb = ONE_DARK_GREEN
    if unit:
        run2 = p.add_run()
        run2.text = unit
        run2.font.name = FONT_BODY
        run2.font.size = Pt(20)
        run2.font.color.rgb = ONE_DARK_GRAY
    p.alignment = PP_ALIGN.LEFT

    # ラベル
    lbl = slide.shapes.add_textbox(
        left + Emu(80000), top + height - Emu(400000),
        width - Emu(160000), Emu(300000)
    )
    lp = lbl.text_frame.paragraphs[0]
    lp.text = label
    lp.font.name = FONT_BODY
    lp.font.size = Pt(16)
    lp.font.color.rgb = ONE_DARK_GRAY

    # 変化率
    if delta:
        dbx = slide.shapes.add_textbox(
            left + Emu(80000), top + Emu(800000),
            Emu(800000), Emu(250000)
        )
        dp = dbx.text_frame.paragraphs[0]
        dp.text = delta
        dp.font.name = FONT_BODY
        dp.font.size = Pt(14)
        dp.font.color.rgb = ONE_DARK_GREEN

    return shape


def build_kpi_row_slide(prs, title, kpis, insight="", source=""):
    """
    KPI-Rowパターン
    kpis: [{"number": "1,200万", "label": "MAU", "unit": "人", "delta": "▲+15%"}, ...]
    """
    slide = add_blank_slide(prs)
    add_title(slide, title)
    add_green_accent_line(slide)

    n = len(kpis)
    gap = 150000
    card_w = int((8128000 - (n - 1) * gap) / n)
    card_h = 1500000

    for i, kpi in enumerate(kpis):
        x = int(MARGIN_LEFT) + i * (card_w + gap)
        _add_kpi_card(
            slide, Emu(x), CONTENT_TOP, Emu(card_w), Emu(card_h),
            kpi["number"], kpi["label"],
            kpi.get("unit", ""), kpi.get("delta", "")
        )

    if insight:
        txBox = slide.shapes.add_textbox(
            MARGIN_LEFT, Emu(2500000), CONTENT_WIDTH, Emu(1800000)
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = insight
        p.font.name = FONT_BODY
        p.font.size = Pt(18)
        p.font.color.rgb = BLACK

    if source:
        add_footnote(slide, source)

    return slide


# ============================================================
# パターン: テーブル
# ============================================================

def build_table_slide(prs, title, headers, rows, highlight_cells=None, source=""):
    """Full-Width Tableパターン"""
    slide = add_blank_slide(prs)
    add_title(slide, title)
    add_green_accent_line(slide)

    highlight_cells = highlight_cells or set()
    n_rows = len(rows) + 1
    n_cols = len(headers)
    table_height = min(n_rows * 500000, 3600000)

    tbl_shape = slide.shapes.add_table(
        n_rows, n_cols,
        MARGIN_LEFT, CONTENT_TOP,
        CONTENT_WIDTH, Emu(table_height)
    )
    table = tbl_shape.table

    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLACK
        for p in cell.text_frame.paragraphs:
            p.font.name = FONT_BODY
            p.font.size = Pt(18)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.alignment = PP_ALIGN.CENTER

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = str(val)
            if (i, j) in highlight_cells:
                cell.fill.solid()
                cell.fill.fore_color.rgb = ONE_LIGHT_GREEN
            elif i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = ONE_LIGHT_GRAY
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT_BODY
                p.font.size = Pt(18)
                p.font.color.rgb = BLACK

    if source:
        add_footnote(slide, source)

    return slide


# ============================================================
# パターン: チャート
# ============================================================

def build_bar_chart_slide(prs, title, categories, series_data, insight="", source=""):
    """棒グラフスライド"""
    slide = add_blank_slide(prs)
    add_title(slide, title)
    add_green_accent_line(slide)

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for s in series_data:
        chart_data.add_series(s["name"], s["values"])

    chart_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        MARGIN_LEFT, CONTENT_TOP,
        CONTENT_WIDTH, Emu(3200000),
        chart_data
    )
    chart = chart_frame.chart
    for i, series in enumerate(chart.series):
        fill = series.format.fill
        fill.solid()
        fill.fore_color.rgb = CHART_COLORS[i % len(CHART_COLORS)]

    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = ONE_LIGHT_GRAY

    if insight:
        txBox = slide.shapes.add_textbox(
            MARGIN_LEFT, Emu(4100000), CONTENT_WIDTH, Emu(400000)
        )
        p = txBox.text_frame.paragraphs[0]
        p.text = insight
        p.font.name = FONT_BODY
        p.font.size = Pt(18)
        p.font.color.rgb = BLACK

    if source:
        add_footnote(slide, source)

    return slide


def build_line_chart_slide(prs, title, categories, series_data, insight="", source=""):
    """折れ線グラフスライド"""
    slide = add_blank_slide(prs)
    add_title(slide, title)
    add_green_accent_line(slide)

    chart_data = CategoryChartData()
    chart_data.categories = categories
    for s in series_data:
        chart_data.add_series(s["name"], s["values"])

    chart_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE_MARKERS,
        MARGIN_LEFT, CONTENT_TOP,
        CONTENT_WIDTH, Emu(3200000),
        chart_data
    )
    chart = chart_frame.chart
    for i, series in enumerate(chart.series):
        series.format.line.color.rgb = CHART_COLORS[i % len(CHART_COLORS)]
        series.format.line.width = Pt(3)

    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = ONE_LIGHT_GRAY

    if insight:
        txBox = slide.shapes.add_textbox(
            MARGIN_LEFT, Emu(4100000), CONTENT_WIDTH, Emu(400000)
        )
        p = txBox.text_frame.paragraphs[0]
        p.text = insight
        p.font.name = FONT_BODY
        p.font.size = Pt(18)
        p.font.color.rgb = BLACK

    if source:
        add_footnote(slide, source)

    return slide


# ============================================================
# パターン: プロセスフロー
# ============================================================

def build_process_flow_slide(prs, title, steps, key_step_index=None, source=""):
    """プロセスフローパターン"""
    slide = add_blank_slide(prs)
    add_title(slide, title)
    add_green_accent_line(slide)

    n = len(steps)
    gap = 400000
    node_w = int((8128000 - (n - 1) * gap) / n)
    node_h = 800000
    y = 2000000

    for i, step in enumerate(steps):
        x = int(MARGIN_LEFT) + i * (node_w + gap)
        is_key = (i == key_step_index)

        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Emu(x), Emu(y), Emu(node_w), Emu(node_h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = ONE_DARK_GREEN if is_key else ONE_LIGHT_GRAY
        shape.line.fill.background()
        shape.adjustments[0] = 0.05

        tf = shape.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        p = tf.paragraphs[0]
        p.text = step
        p.font.name = FONT_BODY
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = WHITE if is_key else BLACK

        if i < n - 1:
            arrow_x = x + node_w + 50000
            arrow_y = y + 350000
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Emu(arrow_x), Emu(arrow_y),
                Emu(300000), Emu(100000)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = ONE_DARK_GRAY
            arrow.line.fill.background()

    if source:
        add_footnote(slide, source)

    return slide


# ============================================================
# パターン: 比較 50/50
# ============================================================

def build_comparison_slide(prs, title, left_header, right_header,
                           left_items, right_items, source=""):
    """比較 50/50 パターン"""
    slide = add_blank_slide(prs)
    add_title(slide, title)
    add_green_accent_line(slide)

    col_w = 3800000
    left_x = int(MARGIN_LEFT)
    right_x = 4836000

    # ヘッダー
    for x, header, color in [
        (left_x, left_header, BLACK),
        (right_x, right_header, ONE_DARK_GREEN)
    ]:
        tx = slide.shapes.add_textbox(
            Emu(x), CONTENT_TOP, Emu(col_w), Emu(350000)
        )
        p = tx.text_frame.paragraphs[0]
        p.text = header
        p.font.name = FONT_TITLE
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = color

    # 区切り線
    div_x = 4568000
    line = slide.shapes.add_connector(
        1, Emu(div_x), CONTENT_TOP, Emu(div_x), Emu(4400000)
    )
    line.line.color.rgb = ONE_DARK_GRAY
    line.line.width = Pt(2)

    # コンテンツ
    for x, items in [(left_x, left_items), (right_x, right_items)]:
        for i, item in enumerate(items):
            ty = int(CONTENT_TOP) + 450000 + i * 500000
            tx = slide.shapes.add_textbox(
                Emu(x), Emu(ty), Emu(col_w), Emu(400000)
            )
            tf = tx.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = f"• {item}"
            p.font.name = FONT_BODY
            p.font.size = Pt(18)
            p.font.color.rgb = BLACK

    if source:
        add_footnote(slide, source)

    return slide


# ============================================================
# パターン: ハイライトボックス
# ============================================================

def add_highlight_box(slide, text):
    """ダークグレーのハイライトボックス"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Emu(1500000), Emu(3500000), Emu(6144000), Emu(600000)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ONE_DARK_GRAY
    shape.line.fill.background()
    shape.adjustments[0] = 0.05

    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_BODY
    p.font.size = Pt(22)
    p.font.color.rgb = WHITE
    return shape


# ============================================================
# デザインチェック
# ============================================================

VALID_COLORS = {
    "3EE3A6", "00D184", "D7FFF0", "000000", "F0F3F6", "8C909B", "FFFFFF"
}


def validate_deck(prs, skip_template_slides=5):
    """デッキのデザインルール準拠をチェック（テンプレートスライドはスキップ）"""
    issues = []

    for i, slide in enumerate(prs.slides, 1):
        if i <= skip_template_slides:
            continue
        bg_color = None
        try:
            fill = slide.background.fill
            if fill.type is not None:
                bg_color = str(fill.fore_color.rgb)
        except Exception:
            pass

        # 黒背景チェック
        if bg_color == "000000":
            issues.append(f"Slide {i}: #000000を背景に使用（禁止）")

        # 緑背景チェック
        if bg_color in ("3EE3A6", "00D184"):
            issues.append(f"Slide {i}: グリーンを背景に使用（禁止）")

        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        size = run.font.size
                        if size and size < Pt(14):
                            issues.append(
                                f"Slide {i}: フォントサイズ {size} < 14pt（禁止）"
                            )

    if not issues:
        return "✅ デザインチェック: すべてのルールに準拠"
    return "⚠️ デザインチェック:\n" + "\n".join(f"  - {x}" for x in issues)


# ============================================================
# メイン（CLIエントリポイント）
# ============================================================

def main():
    parser = argparse.ArgumentParser(description="ONE Slide Builder")
    parser.add_argument("--template", required=True, help="テンプレートpptxパス")
    parser.add_argument("--output", default="output.pptx", help="出力pptxパス")
    parser.add_argument("--config", help="スライド構成JSONファイル（任意）")
    args = parser.parse_args()

    prs = Presentation(args.template)

    if args.config:
        with open(args.config) as f:
            config = json.load(f)
        # JSON設定に基づいてスライドを生成
        for slide_def in config.get("slides", []):
            pattern = slide_def["pattern"]
            if pattern == "statement":
                build_statement_slide(prs, slide_def["statement"],
                                      slide_def.get("sub_text", ""))
            elif pattern == "kpi_row":
                build_kpi_row_slide(prs, slide_def["title"], slide_def["kpis"],
                                    slide_def.get("insight", ""),
                                    slide_def.get("source", ""))
            elif pattern == "table":
                build_table_slide(prs, slide_def["title"],
                                  slide_def["headers"], slide_def["rows"],
                                  set(tuple(c) for c in slide_def.get("highlight_cells", [])),
                                  slide_def.get("source", ""))
            elif pattern == "bar_chart":
                build_bar_chart_slide(prs, slide_def["title"],
                                      slide_def["categories"],
                                      slide_def["series"],
                                      slide_def.get("insight", ""),
                                      slide_def.get("source", ""))
            elif pattern == "line_chart":
                build_line_chart_slide(prs, slide_def["title"],
                                       slide_def["categories"],
                                       slide_def["series"],
                                       slide_def.get("insight", ""),
                                       slide_def.get("source", ""))
            elif pattern == "process_flow":
                build_process_flow_slide(prs, slide_def["title"],
                                         slide_def["steps"],
                                         slide_def.get("key_step_index"),
                                         slide_def.get("source", ""))
            elif pattern == "comparison":
                build_comparison_slide(prs, slide_def["title"],
                                       slide_def["left_header"],
                                       slide_def["right_header"],
                                       slide_def["left_items"],
                                       slide_def["right_items"],
                                       slide_def.get("source", ""))
    else:
        print("--config が指定されていません。テンプレートのみ出力します。")

    # デザインチェック
    result = validate_deck(prs)
    print(result)

    prs.save(args.output)
    print(f"\n保存完了: {args.output}")
    print(f"スライド数: {len(prs.slides)}")


if __name__ == "__main__":
    main()
